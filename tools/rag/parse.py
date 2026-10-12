"""RFS 챗봇 1단계: 원본 자료 → 텍스트 (docs.jsonl)

입력
  - 카페 자료: 비공개 저장소 rfs-cafe-data (글 폴더마다 meta.json, body.txt, 첨부파일)
  - 사이트 리포트: 이 저장소 assets/reports/*.pdf
출력 (카페 자료 저장소 안 parsed/ 에 쓴다. 공개 저장소에는 텍스트를 커밋하지 않는다)
  - docs.jsonl   문서 1개당 1줄. 텍스트, 페이지, 출처, 문서 종류, 중복 여부
  - stats.json   형식·종류별 개수, 글자 수, 중복·실패 목록

  python3 -I tools/rag/parse.py /home/user/rfs-cafe-data
"""
import hashlib
import json
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

import olefile
import openpyxl
import pymupdf
from datasketch import MinHash, MinHashLSH

CAFE = Path(sys.argv[1] if len(sys.argv) > 1 else '/home/user/rfs-cafe-data')
SITE = Path(__file__).resolve().parents[2] / 'assets' / 'reports'
OUT = CAFE / 'parsed'
TMP = Path(tempfile.mkdtemp(prefix='rfs-parse-'))

# ---------------------------------------------------------------- 형식별 추출

def clean_pages(pages):
    """PDF 페이지 텍스트에서 머리말·꼬리말(절반 넘는 페이지에 반복되는 줄)과 차트 축 숫자를 지운다."""
    lines_per_page = [[l.strip() for l in p.splitlines()] for p in pages]
    seen = Counter(l for ls in lines_per_page for l in set(ls) if l)
    n = len(pages)
    repeated = {l for l, c in seen.items() if n >= 4 and c > n * 0.5}
    numeric = re.compile(r'^[\d\s.,%()\-+~/:FEAQ년월원배x×$₩]*$')
    out = []
    for ls in lines_per_page:
        keep = [l for l in ls if l and l not in repeated and not (len(l) <= 14 and numeric.match(l))]
        out.append('\n'.join(keep))
    return out, len(repeated)


def pdf(path):
    doc = pymupdf.open(path)
    raw = [p.get_text() for p in doc]
    pages, headers = clean_pages(raw)
    low = sum(1 for p in raw if len(p.strip()) < 50)
    return {'pages': pages, 'n_pages': len(raw), 'low_text_pages': low, 'removed_header_lines': headers}


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def docx(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read('word/document.xml'))
    out = []
    for el in root.iter():
        if el.tag == W + 'p':
            t = ''.join(x.text or '' for x in el.iter(W + 't')).strip()
            if t:
                out.append(t)
    return {'text': '\n'.join(out)}


def pptx(path):
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)),
                       key=lambda n: int(re.findall(r'\d+', n)[-1]))
        pages = []
        for n in names:
            root = ET.fromstring(z.read(n))
            paras = [''.join(x.text or '' for x in p.iter(A + 't')).strip() for p in root.iter(A + 'p')]
            pages.append('\n'.join(p for p in paras if p))
    return {'pages': pages, 'n_pages': len(pages)}


def hwp(path):
    """HWP 5.0 본문 추출. BodyText/Section* 의 PARA_TEXT(태그 67) 레코드만 읽는다."""
    f = olefile.OleFileIO(str(path))
    compressed = f.openstream('FileHeader').read()[36] & 1
    secs = sorted((s for s in f.listdir() if s[0] == 'BodyText'), key=lambda s: int(s[1][7:] or 0))
    out = []
    for s in secs:
        data = f.openstream(s).read()
        if compressed:
            data = zlib.decompress(data, -15)
        i = 0
        while i + 4 <= len(data):
            h = struct.unpack_from('<I', data, i)[0]
            tag, size = h & 0x3FF, (h >> 20) & 0xFFF
            i += 4
            if size == 0xFFF:
                size = struct.unpack_from('<I', data, i)[0]
                i += 4
            if tag == 67:
                out.append(_hwp_text(data[i:i + size]))
            i += size
    return {'text': '\n'.join(t for t in out if t.strip())}


def _hwp_text(b):
    chars, i = [], 0
    while i + 2 <= len(b):
        c = struct.unpack_from('<H', b, i)[0]
        if c in (1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23):
            i += 16  # 확장·인라인 컨트롤은 8글자 크기
            if c == 9:
                chars.append('\t')
            continue
        if c in (10, 13):
            chars.append('\n')
        elif c >= 32:
            chars.append(chr(c))
        i += 2
    return ''.join(chars).strip()


# 엑셀: 밸류에이션·가정·추정 시트는 수식을 항목 이름으로 풀어서 문장처럼, 나머지 시트는 구조만
KEY_SHEET = re.compile(r'valu|wacc|assum|가정|추정|estim|dcf|per\b|pbr|ev|target|목표|peer|summary|요약|beta|sotp|rim|매출|비용|cost|revenue|sales', re.I)
REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_가-힣]+))?!?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")


def _fmt(v):
    if isinstance(v, float):
        return f'{v:,.4g}' if abs(v) < 1e6 else f'{v:,.0f}'
    return str(v)


def xlsx(path):
    wf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grids, labels = {}, {}
    for ws in wf.worksheets:
        g = {}
        for row in ws.iter_rows(max_row=250, max_col=40):
            for c in row:
                if getattr(c, 'value', None) is not None:
                    g[(c.row, c.column)] = c.value
        grids[ws.title] = g
        rl = {}
        for (r, col), v in sorted(g.items()):
            if r not in rl and isinstance(v, str) and not v.startswith('=') and v.strip():
                rl[r] = v.strip()
        labels[ws.title] = rl
    vals = {ws.title: {(c.row, c.column): c.value for row in ws.iter_rows(max_row=250, max_col=40) for c in row
                       if getattr(c, 'value', None) is not None} for ws in wv.worksheets}

    def col_n(s):
        n = 0
        for ch in s:
            n = n * 26 + ord(ch) - 64
        return n

    def label_of(sheet, r):
        return labels.get(sheet, {}).get(r)

    def explain(sheet, f):
        def sub(m):
            sh = m.group(1) or m.group(2) or sheet
            if sh not in labels:
                sh = sheet
            r1, r2 = int(m.group(4)), int(m.group(6) or 0)
            a = label_of(sh, r1)
            if not a:
                return m.group(0)
            if r2 and r2 != r1:
                b = label_of(sh, r2) or m.group(0)
                return f'[{a}~{b}]'
            return f'[{a}]'
        return REF.sub(sub, f[1:])

    parts = []
    for name, g in grids.items():
        if not g:
            continue
        if KEY_SHEET.search(name):
            lines = []
            rows = defaultdict(list)
            for (r, c), v in sorted(g.items()):
                rows[r].append((c, v))
            for r, cells in rows.items():
                out = []
                for c, v in cells:
                    cv = vals.get(name, {}).get((r, c))
                    if isinstance(v, str) and v.startswith('='):
                        shown = _fmt(cv) if cv is not None else '?'
                        out.append(f'{shown} (= {explain(name, v)[:120]})')
                    else:
                        out.append(_fmt(v))
                line = ' | '.join(out)
                if len(line) > 2:
                    lines.append(line[:400])
            parts.append(f'## 시트: {name}\n' + '\n'.join(lines[:200]))
        else:
            heads = sorted({_fmt(v) for (r, c), v in g.items() if r <= 6 and v is not None})[:30]
            items = list(dict.fromkeys(labels[name].values()))[:80]
            parts.append(f'## 시트: {name} (구조만)\n열: {", ".join(heads)}\n항목: {", ".join(items)}')
    return {'text': '\n\n'.join(parts), 'sheets': list(grids)}


def convert(path, fmt):
    """ppt·xls·doc → LibreOffice로 새 형식 변환"""
    d = TMP / hashlib.md5(str(path).encode()).hexdigest()
    d.mkdir(exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', str(d), str(path)],
                   capture_output=True, timeout=180)
    got = list(d.glob('*.' + fmt))
    if not got:
        raise RuntimeError('변환 실패')
    return got[0]


def extract(path):
    ext = path.suffix.lower()
    if ext == '.pdf':
        return pdf(path)
    if ext == '.docx':
        return docx(path)
    if ext == '.pptx':
        return pptx(path)
    if ext in ('.xlsx', '.xlsm'):
        return xlsx(path)
    if ext == '.hwp':
        return hwp(path)
    if ext == '.ppt':
        return pptx(convert(path, 'pptx'))
    if ext == '.xls':
        return xlsx(convert(path, 'xlsx'))
    if ext == '.doc':
        return docx(convert(path, 'docx'))
    if ext in ('.txt', '.csv'):
        return {'text': path.read_text(errors='ignore')}
    raise ValueError('지원 안 함 ' + ext)


# ---------------------------------------------------------------- 문서 종류

def doc_type(board, group, subject, fname, ext):
    s = f'{subject} {fname}'
    if ext in ('.xlsx', '.xlsm', '.xls'):
        return 'model'
    if board == '투자제안서' or '투자제안서' in s:
        return 'proposal'
    if re.search(r'가이드|guide|양식|템플릿|template|작성법|작성 요령|교육', s, re.I) or board in ('R.F.S 가이드북',):
        return 'guide'
    if board in ('운용 현황 보고',):
        return 'fund_report'
    if board in ('매크로 보고서', '22 매크로 이슈 브리핑'):
        return 'macro'
    if re.search(r'산업\s*(분석|리포트|보고서)|industry|산업$', s, re.I) or board == '21 산업 분석':
        return 'industry'
    if group == '소모임' or group == 'RFS EDUCATION':
        return 'study'
    return 'company'


GRADE = {'company': 'A', 'industry': 'A', 'guide': 'A', 'model': 'A', 'macro': 'B', 'study': 'B',
         'fund_report': 'B', 'post': 'B', 'proposal': 'C'}

# ---------------------------------------------------------------- 수집

def files_of(adir):
    for p in sorted(adir.iterdir()):
        if p.name in ('meta.json', 'body.txt') or p.is_dir():
            continue
        if p.suffix.lower() == '.zip':
            d = TMP / ('zip-' + hashlib.md5(str(p).encode()).hexdigest())
            try:
                with zipfile.ZipFile(p) as z:
                    for info in z.infolist():
                        name = info.filename
                        if not (info.flag_bits & 0x800):  # 윈도우 압축은 cp949 이름
                            try:
                                name = name.encode('cp437').decode('cp949')
                            except Exception:
                                pass
                        if info.is_dir() or '__MACOSX' in name:
                            continue
                        tgt = d / Path(name).name
                        d.mkdir(parents=True, exist_ok=True)
                        tgt.write_bytes(z.read(info))
                        yield tgt, p.name + '/' + Path(name).name
            except Exception as e:
                yield p, p.name  # 아래에서 실패로 기록
            continue
        yield p, p.name


def main():
    OUT.mkdir(exist_ok=True)
    docs, fails = [], []
    arts = sorted(CAFE.glob('*/*/*/meta.json'))
    for mp in arts:
        meta = json.loads(mp.read_text())
        adir = mp.parent
        base = {'source': 'cafe', 'group': meta['group'], 'board': meta['board'], 'article_id': meta['articleId'],
                'subject': meta['subject'], 'date': meta.get('date')}
        body = adir / 'body.txt'
        if body.exists() and len(body.read_text().strip()) >= 200:
            docs.append({**base, 'file': None, 'ext': '.post', 'doc_type': 'post', 'text': body.read_text().strip()})
        for p, shown in files_of(adir):
            ext = p.suffix.lower()
            if ext == '.zip':
                fails.append({'file': str(p.relative_to(CAFE)), 'error': 'zip 열기 실패'})
                continue
            try:
                r = extract(p)
            except Exception as e:
                fails.append({'file': f'{adir.relative_to(CAFE)}/{shown}', 'error': str(e)[:120]})
                continue
            t = r.get('text') or '\n\n'.join(r.get('pages', []))
            dt = doc_type(meta['board'], meta['group'], meta['subject'], shown, ext)
            docs.append({**base, 'file': shown, 'ext': ext, 'doc_type': dt, 'text': t,
                         **{k: v for k, v in r.items() if k != 'text'}})
    for p in sorted(SITE.glob('*.pdf')):
        r = pdf(p)
        sem, slug = p.stem.split('_', 1)
        docs.append({'source': 'site', 'group': '사이트 리포트', 'board': '리포트', 'article_id': None,
                     'subject': slug, 'semester': sem, 'date': None, 'file': p.name, 'ext': '.pdf',
                     'doc_type': 'company', 'text': '\n\n'.join(r['pages']), **r})

    for i, d in enumerate(docs):
        d['doc_id'] = f'd{i:05d}'
        d['grade'] = GRADE[d['doc_type']]
        d['chars'] = len(d['text'])

    # 중복: 같은 리포트를 PDF·워드로 둘 다 올렸거나, 사이트와 카페에 같은 리포트가 있는 경우
    def norm(t):
        return re.sub(r'\s+', '', t)[:20000]

    lsh = MinHashLSH(threshold=0.8, num_perm=128)
    mh = {}
    for d in docs:
        n = norm(d['text'])
        if len(n) < 300:
            continue
        m = MinHash(num_perm=128)
        for k in range(0, len(n) - 5, 2):
            m.update(n[k:k + 5].encode())
        mh[d['doc_id']] = m
        lsh.insert(d['doc_id'], m)
    by_id = {d['doc_id']: d for d in docs}
    prio = lambda d: (d['source'] == 'site', d['ext'] == '.pdf', d['ext'] == '.pptx', d['chars'])
    seen = set()
    groups = []
    for did, m in mh.items():
        if did in seen:
            continue
        g = [x for x in lsh.query(m) if x not in seen]
        seen.update(g)
        if len(g) > 1:
            g.sort(key=lambda x: prio(by_id[x]), reverse=True)
            groups.append(g)
            for x in g[1:]:
                by_id[x]['duplicate_of'] = g[0]

    with open(OUT / 'docs.jsonl', 'w') as f:
        for d in docs:
            f.write(json.dumps(d, ensure_ascii=False) + '\n')

    live = [d for d in docs if 'duplicate_of' not in d]
    empty = [d for d in live if d['chars'] < 200]
    stats = {
        'articles': len(arts),
        'docs_total': len(docs),
        'docs_after_dedup': len(live),
        'duplicates_removed': len(docs) - len(live),
        'dup_groups': len(groups),
        'site_reports_dup_in_cafe': sum(1 for g in groups if any(by_id[x]['source'] == 'site' for x in g)
                                        and any(by_id[x]['source'] == 'cafe' for x in g)),
        'chars_after_dedup': sum(d['chars'] for d in live),
        'by_ext': Counter(d['ext'] for d in live),
        'by_type': Counter(d['doc_type'] for d in live),
        'by_grade': Counter(d['grade'] for d in live),
        'chars_by_type': {k: sum(d['chars'] for d in live if d['doc_type'] == k) for k in GRADE},
        'pdf_low_text_pages': sum(d.get('low_text_pages', 0) for d in live),
        'pdf_pages': sum(d.get('n_pages', 0) for d in live if d['ext'] == '.pdf'),
        'near_empty_docs': [{'file': d['file'], 'subject': d['subject'], 'chars': d['chars']} for d in empty][:200],
        'near_empty_count': len(empty),
        'failures': fails,
    }
    (OUT / 'stats.json').write_text(json.dumps(stats, ensure_ascii=False, indent=1))
    shutil.rmtree(TMP, ignore_errors=True)
    print(json.dumps({k: v for k, v in stats.items() if k not in ('near_empty_docs', 'failures')}, ensure_ascii=False, indent=1))
    print('실패', len(fails))


if __name__ == '__main__':
    main()
